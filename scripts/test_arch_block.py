# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_native_arch_diagram(doc, caption_prefix, caption_title):
    table = doc.add_table(rows=5, cols=1)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    width_in = 6.2
    
    tiers = [
        ("TẦNG THIẾT BỊ MÁY TRẠM TẠI QUẦY (CLIENT STATIONS)",
         "• Màn hình khách (/display): Camera 30 fps, Web Worker MediaPipe bám 21 mốc tay + 11 mốc thân, cắt chuẩn hóa 16 khung hình RGB 224x224 (trim-bodycrop-v1).\n• Bàn làm việc nhân viên (/recognize): Nhận video WebRTC P2P nội bộ, bộ ghép câu rảnh tay (Sentence Composer), phát âm TTS và micro phản hồi STT."),
        
        ("TẦNG MÁY CHỦ DỊCH VỤ & ĐIỀU PHỐI (FASTAPI BACKEND)",
         "• FastAPI Gateway & Uvicorn ASGI Server: Quản lý phiên làm việc đa quầy, điều phối WebSocket Hub.\n• Ghép đôi quầy an toàn (PIN 6 số, QR Token, Cookie Kiosk vĩnh viễn), xác thực RBAC & băm Argon2id.\n• Bộ điều phối luồng (Debounce 1.000 ms, Heuristic nghỉ 1.800 ms), ký số Webhook HMAC SHA-256.\n• Bộ chuyển mạch dự phòng Failover: Tự động điều tuyến giữa Asia Modal GPU và Local Core ML trong 6s."),
         
        ("TẦNG HẠ TẦNG SUY LUẬN AI MÁY CHỦ (SERVER-SIDE GPU)",
         "• Cụm GPU đám mây Đông Nam Á (Asia Modal GPU): NVIDIA L4 24GB Tensor Core, PyTorch CUDA 12, thực thi mô hình video không-thời gian MViTv2-S (sales30-epoch36), thời gian tính toán 35–65 ms / 16 frame.\n• Máy chủ trạm dự phòng tại chỗ: Apple Silicon Neural Engine (Core ML port 8000), suy luận < 15 ms.\n• Bảo mật riêng tư: 16 khung hình RGB chỉ giải mã trong RAM GPU của tiến trình, tuyệt đối không ghi đĩa.")
    ]
    
    connectors = [
        "▲   Kênh kết nối thời gian thực: WebSocket (16 frame RGB)  &  WebRTC P2P (Luồng video LAN nội bộ)   ▲\n▼                                                                                                                                           ▼",
        "▲   Giao thức mạng suy luận AI: wss:// (Cloud GPU chính)  &  HTTP localhost:8000 (Dự phòng Core ML)   ▲\n▼                                                                                                                                           ▼"
    ]
    
    row_idx = 0
    for i, (title, content) in enumerate(tiers):
        cell = table.rows[row_idx].cells[0]
        cell.width = Inches(width_in)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
            f'<w:left w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
            f'<w:right w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        tcPr.append(shd)
        
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="120" w:type="dxa"/>'
            f'<w:bottom w:w="120" w:type="dxa"/>'
            f'<w:left w:w="180" w:type="dxa"/>'
            f'<w:right w:w="180" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)
        
        p = cell.paragraphs[0]
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(title + "\n")
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.bold = True
        r_title.font.color.rgb = RGBColor(0, 0, 0)
        
        p2 = cell.add_paragraph()
        p2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = 1.15
        r_content = p2.add_run(content)
        r_content.font.name = "Times New Roman"
        r_content.font.size = Pt(11)
        r_content.font.color.rgb = RGBColor(30, 41, 59)
        
        if i < len(tiers) - 1:
            conn_cell = table.rows[row_idx + 1].cells[0]
            conn_cell.width = Inches(width_in)
            tcPr_c = conn_cell._tc.get_or_add_tcPr()
            borders_c = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="none"/>'
                f'<w:left w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr_c.append(borders_c)
            p_c = conn_cell.paragraphs[0]
            p_c.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.space_before = Pt(2)
            p_c.paragraph_format.space_after = Pt(2)
            r_c = p_c.add_run(connectors[i])
            r_c.font.name = "Times New Roman"
            r_c.font.size = Pt(10)
            r_c.bold = True
            r_c.font.color.rgb = RGBColor(0, 0, 0)
            row_idx += 2
        else:
            row_idx += 1
            
    p_cap = doc.add_paragraph()
    p_cap.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(14)
    r_pre = p_cap.add_run(caption_prefix + ": ")
    r_pre.font.name = "Times New Roman"
    r_pre.font.size = Pt(12)
    r_pre.bold = True
    r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_title = p_cap.add_run(caption_title)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(12)
    r_title.italic = True
    r_title.font.color.rgb = RGBColor(51, 65, 85)

doc = docx.Document()
create_native_arch_diagram(doc, "Sơ đồ 5.1", "Sơ đồ khối kiến trúc phân tầng hệ thống VSign")
doc.save("/tmp/test_arch_block.docx")
print("Saved /tmp/test_arch_block.docx successfully!")
