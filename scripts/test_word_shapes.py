# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml

doc = docx.Document()

def add_native_shape_box(doc, title, subtitle="", width_emu=4500000, height_emu=800000):
    p = doc.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    sub_xml = ""
    if subtitle:
        sub_xml = (
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:pPr><w:jc w:val="center"/><w:spacing w:before="0" w:after="40"/></w:pPr>'
            '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            '<w:sz w:val="20"/><w:color val="444444"/></w:rPr>'
            f'<w:t>{subtitle}</w:t></w:r></w:p>'
        )

    xml = f"""
    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
         xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
         xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
         xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <w:drawing>
        <wp:inline distT="0" distB="0" distL="0" distR="0">
          <wp:extent cx="{width_emu}" cy="{height_emu}"/>
          <wp:docPr id="101" name="WordShape"/>
          <wps:wsp>
            <wps:cNvSpPr/>
            <wps:spPr>
              <a:xfrm>
                <a:off x="0" y="0"/>
                <a:ext cx="{width_emu}" cy="{height_emu}"/>
              </a:xfrm>
              <a:prstGeom prst="roundRect">
                <a:avLst><a:gd name="adj" fmla="val 6000"/></a:avLst>
              </a:prstGeom>
              <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
              <a:ln w="19050">
                <a:solidFill><a:srgbClr val="000000"/></a:solidFill>
              </a:ln>
            </wps:spPr>
            <wps:txbx>
              <w:txbxContent>
                <w:p>
                  <w:pPr><w:jc w:val="center"/><w:spacing w:before="60" w:after="20"/></w:pPr>
                  <w:r>
                    <w:rPr>
                      <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                      <w:sz w:val="26"/><w:b/><w:color val="000000"/>
                    </w:rPr>
                    <w:t>{title}</w:t>
                  </w:r>
                </w:p>
                {sub_xml}
              </w:txbxContent>
            </wps:txbx>
          </wps:wsp>
        </wp:inline>
      </w:drawing>
    </w:r>
    """
    p._p.append(parse_xml(xml))

def add_native_shape_arrow(doc):
    p = doc.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("▼")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

add_native_shape_box(doc, "1. THU NHẬN HÌNH ẢNH", "Camera quầy khách thu nhận video 30 fps")
add_native_shape_arrow(doc)
add_native_shape_box(doc, "2. BÁM BẮT ĐẶC TRƯNG", "Web Worker: MediaPipe 21 mốc tay + 11 mốc thân")
add_native_shape_arrow(doc)
add_native_shape_box(doc, "3. CẮT VÙNG THÂN TRÊN ROI", "Thuật toán trim-bodycrop-v1: 16 khung hình RGB (224x224)")
add_native_shape_arrow(doc)
add_native_shape_box(doc, "4. SUY LUẬN NHẬN DẠNG AI", "MViTv2-S Spatiotemporal (Asia Modal GPU / Local Core ML)")
add_native_shape_arrow(doc)
add_native_shape_box(doc, "5. GHÉP CÂU & PHẢN HỒI", "Sentence Composer rảnh tay + Web Speech TTS đọc to")

doc.save("/tmp/test_flow_drawingml.docx")
print("Saved /tmp/test_flow_drawingml.docx successfully!")
